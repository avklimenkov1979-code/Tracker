#!/usr/bin/env python3
"""Сбор биржевых цен на фенол из бюллетеней SPIMEX (Петербургская биржа).

Скрипт открывает страницы «Итоги торгов» нужных секций, находит ссылки на
ежедневные бюллетени (.xls/.xlsx), скачивает те, что ещё не обработаны,
вытаскивает строки с фенолом и дописывает их в таблицу:

  * data/phenol_prices.csv  — основная таблица (UTF-8 с BOM, разделитель «;»,
    открывается в Excel двойным кликом);
  * data/phenol_prices.xlsx — та же таблица в формате Excel.

Повторный запуск безопасен: строки дедуплицируются по (дата, код инструмента).

Примеры:
  python spimex_phenol.py                 # последний бюллетень каждой секции
  python spimex_phenol.py --days 30       # догрузить историю за 30 дней
  python spimex_phenol.py --file bulletin.xls   # разобрать локальный файл
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urljoin

import requests

log = logging.getLogger("spimex_phenol")

BASE_URL = "https://spimex.com"
# Страницы «Итоги торгов». Если фенол торгуется в другой секции — добавьте её
# страницу сюда или передайте через --results-url.
DEFAULT_RESULTS_URLS = [
    f"{BASE_URL}/markets/oil_products/trades/results/",
]
DEFAULT_PATTERN = r"фенол(?![а-яё])"  # «Фенол», но не «фенольная смола»

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CSV = ROOT / "data" / "phenol_prices.csv"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36",
    "Accept-Language": "ru-RU,ru;q=0.9",
}

# Колонки итоговой таблицы: (ключ, заголовок в таблице).
COLUMNS = [
    ("date", "Дата торгов"),
    ("code", "Код инструмента"),
    ("name", "Наименование инструмента"),
    ("basis", "Базис поставки"),
    ("volume_units", "Объём, т"),
    ("volume_rub", "Объём, руб."),
    ("price_change", "Изменение к пред. дню, руб."),
    ("price_min", "Мин. цена, руб./т"),
    ("price_avg", "Средняя цена, руб./т"),
    ("price_max", "Макс. цена, руб./т"),
    ("price_market", "Рыночная цена, руб./т"),
    ("bid_best", "Лучшая цена покупки, руб./т"),
    ("ask_best", "Лучшая цена продажи, руб./т"),
    ("deals", "Кол-во договоров"),
    ("source", "Бюллетень"),
]

# Как опознать колонку бюллетеня по её заголовку (заголовок нормализуется:
# нижний регистр, без пробелов/переносов, «ё»→«е», «обьем»→«объем»).
HEADER_RULES = [
    ("code", ("кодинструмента",)),
    ("name", ("наименованиеинструмента",)),
    ("basis", ("базиспоставки",)),
    ("volume_units", ("объемдоговороввединицах",)),
    ("volume_rub", ("объемдоговоров,руб", "объемдоговороввруб")),
    ("price_change", ("изменениерыночнойцены",)),
    ("price_min", ("минимальнаяцена",)),
    ("price_avg", ("средняяцена",)),
    ("price_max", ("максимальнаяцена",)),
    ("price_market", ("рыночнаяцена",)),
    ("bid_best", ("покупка", "спрос")),
    ("ask_best", ("продажа", "предложение")),
    ("deals", ("количестводоговоров",)),
]

LINK_RE = re.compile(
    r"""href\s*=\s*["']([^"']+?\.xlsx?(?:\?[^"']*)?)["']""",
    re.IGNORECASE,
)
# Сейчас бюллетени на сайте отдаются ссылками вида /files/61309/ (без расширения).
FILES_LINK_RE = re.compile(
    r"""<a\b[^>]*?href\s*=\s*["']([^"']*?/files/\d+/?(?:\?[^"']*)?)["'][^>]*>(.*?)</a>""",
    re.IGNORECASE | re.DOTALL,
)
FILE_DATE_RE = re.compile(r"(20\d{2})(\d{2})(\d{2})\d{0,6}")
TEXT_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(20\d{2})")
SHEET_DATE_RE = re.compile(r"дата\s+торгов[:\s]*(\d{2})\.(\d{2})\.(\d{4})", re.IGNORECASE)
TAG_RE = re.compile(r"<[^>]+>")


@dataclass
class Bulletin:
    url: str
    date: dt.date | None = None
    rows: list[dict] = field(default_factory=list)
    label: str = ""
    is_bulletin: bool = True
    instruments: int = 0


# ---------------------------------------------------------------- загрузка

def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(HEADERS)
    return s


def source_key(url: str) -> str:
    """Короткий идентификатор файла: имя файла или номер из /files/61309/."""
    path = url.split("?", 1)[0].rstrip("/")
    return path.rsplit("/", 1)[-1]


def date_from_filename(url: str) -> dt.date | None:
    m = FILE_DATE_RE.search(source_key(url))
    if not m:
        return None
    try:
        return dt.date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def date_from_text(text: str) -> dt.date | None:
    for m in reversed(list(TEXT_DATE_RE.finditer(text))):
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            continue
    return None


def find_bulletin_links(html: str, page_url: str) -> list[Bulletin]:
    seen, out = set(), []

    def add(href: str, date: dt.date | None, label: str = "") -> None:
        url = urljoin(page_url, href.replace("&amp;", "&"))
        key = url.split("?", 1)[0]
        if key not in seen:
            seen.add(key)
            out.append(Bulletin(url=url, date=date, label=label))

    for href in LINK_RE.findall(html):
        add(href, date_from_filename(href))
    # Бюллетени лежат в /upload/reports/; если такой папки нет — берём .xls с датой в имени
    reports = [b for b in out if "/reports/" in b.url.lower()]
    xls = reports or [b for b in out if b.date]
    if xls:
        return xls

    out, seen = [], set()
    for m in FILES_LINK_RE.finditer(html):
        label = " ".join(TAG_RE.sub(" ", m[2]).split())
        # дата — в тексте ссылки или в тексте перед ней (строка таблицы бюллетеней)
        before = " ".join(TAG_RE.sub(" ", html[max(0, m.start() - 1500):m.start()]).split())
        add(m[1], date_from_text(label) or date_from_text(before[-300:]), label)
    return out


def list_bulletins(session: requests.Session, results_url: str, since: dt.date | None,
                   max_pages: int = 50) -> list[Bulletin]:
    """Ссылки-кандидаты со страницы итогов, от новых к старым (с пагинацией для истории)."""
    found: list[Bulletin] = []
    if since is not None:
        # если даты у ссылок не распознаются, не листаем весь архив
        max_pages = min(max_pages, 2 + (dt.date.today() - since).days // 7)
    for page in range(1, max_pages + 1):
        url = results_url if page == 1 else f"{results_url}?page=page-{page}"
        resp = session.get(url, timeout=60)
        resp.raise_for_status()
        links = find_bulletin_links(resp.text, url)
        if page == 1 and not links:
            explain_empty_page(resp)
        known = {f.url.split("?")[0] for f in found}
        new = [b for b in links if b.url.split("?")[0] not in known]
        if not new:
            break
        found.extend(new)
        if since is None:
            break  # нужен только последний бюллетень — первой страницы достаточно
        dates = [b.date for b in new if b.date]
        if dates and min(dates) < since:
            break
    if since is not None:
        found = [b for b in found if b.date is None or b.date >= since]
    return found


def fetch_excel(session: requests.Session, url: str) -> bytes | None:
    """Скачать файл, если это Excel; PDF и прочее пропускаются после первых байт."""
    with session.get(url, timeout=120, stream=True) as resp:
        resp.raise_for_status()
        chunks = resp.iter_content(64 * 1024)
        head = next(chunks, b"")
        if not (head.startswith(b"\xd0\xcf\x11\xe0") or head.startswith(b"PK\x03\x04")):
            log.debug("%s: не Excel (%r…) — пропускаю", url, head[:8])
            return None
        return head + b"".join(chunks)


def explain_empty_page(resp: requests.Response) -> None:
    """Если ссылок не нашлось — показать, что вообще пришло, чтобы поправить разбор."""
    html = resp.text
    title = re.search(r"<title[^>]*>(.*?)</title>", html, re.IGNORECASE | re.DOTALL)
    hrefs = re.findall(r"""href\s*=\s*["']([^"']+)["']""", html, re.IGNORECASE)
    hint = [h for h in hrefs if re.search(r"xls|upload|report|bullet|pdf|download|file", h, re.IGNORECASE)]
    log.warning("ссылок на .xls нет. Ответ: HTTP %s, %s, %d байт, заголовок: %r",
                resp.status_code, resp.url, len(html),
                " ".join(title[1].split())[:80] if title else None)
    log.warning("всего ссылок: %d, похожих на файлы: %d", len(hrefs), len(hint))
    for h in hint[:8]:
        log.warning("  %s", h[:120])
    if not hint:
        text = re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", html, flags=re.DOTALL | re.IGNORECASE)
        log.warning("начало текста страницы: %s", " ".join(text.split())[:300])


# ---------------------------------------------------------------- разбор

def norm_header(value) -> str:
    s = str(value or "").lower().replace("ё", "е").replace("обьем", "объем")
    return re.sub(r"\s+", "", s)


def to_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if s in ("", "-", "—", "–"):
        return None
    try:
        num = float(s)
    except ValueError:
        return None
    return int(num) if num.is_integer() else num


def read_sheet_rows(content: bytes, url: str) -> list[list]:
    """Все строки первого листа бюллетеня как списки значений."""
    if content[:4] == b"PK\x03\x04" or url.split("?")[0].lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    import xlrd
    book = xlrd.open_workbook(file_contents=content)
    sh = book.sheet_by_index(0)
    return [sh.row_values(i) for i in range(sh.nrows)]


def map_header(row: list) -> dict[str, int] | None:
    cells = [norm_header(c) for c in row]
    if not any("кодинструмента" in c for c in cells):
        return None
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(cells):
        if not cell:
            continue
        for key, needles in HEADER_RULES:
            if key in mapping:
                continue
            if any(n in cell for n in needles):
                # «Рыночная цена» не должна перехватить «Изменение рыночной цены»
                if key == "price_market" and "изменение" in cell:
                    continue
                mapping[key] = idx
                break
    return mapping if {"code", "name"} <= mapping.keys() else None


def parse_bulletin(content: bytes, url: str, pattern: re.Pattern,
                   fallback_date: dt.date | None = None) -> Bulletin:
    rows = read_sheet_rows(content, url)
    trade_date = None
    for row in rows[:15]:
        for cell in row:
            m = SHEET_DATE_RE.search(str(cell or ""))
            if m:
                trade_date = dt.date(int(m[3]), int(m[2]), int(m[1]))
                break
        if trade_date:
            break
    trade_date = trade_date or fallback_date or date_from_filename(url)

    bulletin = Bulletin(url=url, date=trade_date, is_bulletin=False)
    header: dict[str, int] | None = None
    for row in rows:
        h = map_header(row)
        if h:
            header = h  # в бюллетене несколько таблиц, у каждой свой заголовок
            bulletin.is_bulletin = True
            continue
        if header is None:
            continue
        get = lambda k: row[header[k]] if k in header and header[k] < len(row) else None
        name = str(get("name") or "").strip()
        if name and str(get("code") or "").strip():
            bulletin.instruments += 1
        if not name or not pattern.search(name):
            continue
        rec = {
            "date": trade_date.isoformat() if trade_date else "",
            "code": str(get("code") or "").strip(),
            "name": " ".join(name.split()),
            "basis": " ".join(str(get("basis") or "").split()),
            "source": source_key(url),
        }
        for key, _ in COLUMNS:
            if key not in rec:
                rec[key] = to_number(get(key))
        bulletin.rows.append(rec)
    return bulletin


# ---------------------------------------------------------------- таблица

def load_table(path: Path) -> list[dict]:
    if not path.exists():
        return []
    titles = {title: key for key, title in COLUMNS}
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [{titles.get(k, k): v for k, v in r.items()} for r in csv.DictReader(f, delimiter=";")]


def merge_rows(existing: list[dict], new: list[dict]) -> tuple[list[dict], int]:
    key = lambda r: (str(r.get("date", "")), str(r.get("code", "")))
    merged = {key(r): r for r in existing}
    added = sum(1 for r in new if key(r) not in merged)
    for r in new:
        merged[key(r)] = r
    rows = sorted(merged.values(), key=lambda r: (str(r.get("date", "")), str(r.get("code", ""))))
    return rows, added


def fmt(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}".replace(".", ",")
    return str(value)


def save_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([title for _, title in COLUMNS])
        for r in rows:
            w.writerow([fmt(r.get(k)) for k, _ in COLUMNS])


def save_xlsx(path: Path, rows: list[dict]) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Фенол"
    ws.append([title for _, title in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        line = []
        for k, _ in COLUMNS:
            v = r.get(k)
            if k == "date" and v:
                v = dt.date.fromisoformat(str(v))
            elif k not in ("code", "name", "basis", "source"):
                v = to_number(v)
            line.append(v)
        ws.append(line)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "DD.MM.YYYY"
    for col, width in zip("ABCDEFGHIJKLMNO", (12, 16, 48, 36, 10, 16, 14, 14, 14, 14, 14, 14, 14, 10, 30)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ---------------------------------------------------------------- main

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--days", type=int, default=0,
                    help="догрузить бюллетени за последние N дней (по умолчанию — только последний)")
    ap.add_argument("--results-url", action="append",
                    help="страница «Итоги торгов» секции (можно указать несколько раз)")
    ap.add_argument("--pattern", default=DEFAULT_PATTERN,
                    help="регулярное выражение для наименования инструмента (по умолчанию: %(default)s)")
    ap.add_argument("--file", action="append", type=Path,
                    help="разобрать локальный файл бюллетеня вместо загрузки с сайта")
    ap.add_argument("--csv", type=Path, default=DEFAULT_CSV, help="путь к CSV (по умолчанию: %(default)s)")
    ap.add_argument("--no-xlsx", action="store_true", help="не сохранять копию в .xlsx")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s")
    pattern = re.compile(args.pattern, re.IGNORECASE)
    existing = load_table(args.csv)
    # обработанные файлы (в т.ч. PDF и бюллетени без фенола), чтобы не качать их повторно
    seen_path = args.csv.with_suffix(".processed.txt")
    try:
        processed = set(seen_path.read_text(encoding="utf-8").split())
    except OSError:
        processed = set()
    processed |= {r.get("source") for r in existing if r.get("source")}

    collected: list[dict] = []
    errors = 0

    if args.file:
        for p in args.file:
            b = parse_bulletin(p.read_bytes(), p.name, pattern)
            log.info("%s: дата %s, строк с фенолом: %d", p.name, b.date, len(b.rows))
            collected += b.rows
    else:
        session = make_session()
        since = dt.date.today() - dt.timedelta(days=args.days) if args.days else None
        for results_url in args.results_url or DEFAULT_RESULTS_URLS:
            try:
                candidates = list_bulletins(session, results_url, since)
            except requests.RequestException as e:
                log.error("не удалось открыть %s: %s", results_url, e)
                errors += 1
                continue
            log.info("%s: ссылок-кандидатов: %d", results_url, len(candidates))
            found = downloads = 0
            for b in candidates:
                key = source_key(b.url)
                if "skip:" + key in processed:  # PDF или другой документ
                    continue
                if key in processed:
                    if since is None:
                        log.info("последний бюллетень (%s) уже обработан", key)
                        break
                    continue
                if downloads >= (20 if since is None else 40 + args.days * 3):
                    log.warning("слишком много файлов без бюллетеня — останавливаюсь")
                    break
                downloads += 1
                try:
                    content = fetch_excel(session, b.url)
                    parsed = parse_bulletin(content, b.url, pattern, b.date) if content else None
                except Exception as e:  # битый файл не должен ронять весь прогон
                    log.error("ошибка при обработке %s: %s", b.url, e)
                    errors += 1
                    continue
                if parsed is None or not parsed.is_bulletin:
                    log.debug("%s (%s): не бюллетень", key, b.label)
                    processed.add("skip:" + key)
                    continue
                if since is not None and parsed.date and parsed.date < since:
                    break
                found += 1
                processed.add(key)
                log.info("бюллетень %s от %s: инструментов %d, строк с фенолом: %d",
                         key, parsed.date, parsed.instruments, len(parsed.rows))
                collected += parsed.rows
                if since is None:
                    break
            if not found:
                log.warning("%s: ни одного бюллетеня Excel не найдено", results_url)
                for b in candidates[:6]:
                    log.warning("  %s  %s  %s", source_key(b.url), b.date or "-", b.label[:60])

    rows, added = merge_rows(existing, collected)
    if collected or not args.csv.exists():
        save_csv(args.csv, rows)
        if not args.no_xlsx:
            save_xlsx(args.csv.with_suffix(".xlsx"), rows)
    log.info("новых строк: %d, всего в таблице: %d (%s)", added, len(rows), args.csv)
    if not args.file:
        try:
            seen_path.parent.mkdir(parents=True, exist_ok=True)
            seen_path.write_text("\n".join(sorted(processed)) + "\n", encoding="utf-8")
        except OSError as e:
            log.warning("не удалось сохранить %s: %s", seen_path, e)
    if not collected and not errors:
        log.warning("фенол в обработанных бюллетенях не найден "
                    "(нет сделок или инструмент торгуется в другой секции — см. --results-url)")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
