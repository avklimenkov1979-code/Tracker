#!/usr/bin/env python3
"""Ежедневный сбор индикатора цен на фенол (FEN) с сайта Петербургской биржи.

Дневные бюллетени торгов с 28.09.2026 закрыты (Указ №686), но биржа публикует
индикаторы оптовых цен нефтегазохимии, среди них фенол — код FEN. Скрипт
забирает строки FEN со страниц индикаторов и дописывает их в таблицу
(CSV для Excel + копия .xlsx). Один запуск в день — одна строка на каждую
таблицу индикаторов; повторный запуск в тот же день перезаписывает строку.

  python phenol_index.py --csv ~/Documents/SPIMEX/phenol_index.csv
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import logging
import re
import sys
from pathlib import Path

import requests

log = logging.getLogger("phenol_index")

DEFAULT_URLS = ["https://spimex.com/indexes/petrochem/national/"]
DEFAULT_CODE = "FEN"
DEFAULT_CSV = Path(__file__).resolve().parent.parent / "data" / "phenol_index.csv"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X) AppleWebKit/605.1.15 "
    "(KHTML, like Gecko) Version/18.0 Mobile/15E148 Safari/604.1",
    "Accept-Language": "ru-RU,ru;q=0.9",
}

COLUMNS = [
    ("date", "Дата сбора"),
    ("code", "Код"),
    ("name", "Наименование"),
    ("value", "Значение, руб./т"),
    ("change", "Изменение, %"),
    ("title", "Индикатор"),
    ("table", "№ таблицы"),
    ("page", "Страница"),
]

TAG_RE = re.compile(r"<[^>]+>")
ROW_START_RE = re.compile(r"""<div\b[^>]*class=["'][^"']*indexes-table__row""", re.I)
VALUE_RE = re.compile(r"(?<![\d.,])(\d{1,3}(?:[ \xa0 ]\d{3})+(?:[.,]\d+)?|\d{4,}(?:[.,]\d+)?)(?![\d%])")
CHANGE_RE = re.compile(r"([+\-−–]?\s?\d+(?:[.,]\d+)?)\s*%")
TITLE_RE = re.compile(r"(?:Национальн|Территориальн|Сводн)[^()]{0,400}?\([^()]{0,80}\)", re.I)
TITLE_START_RE = re.compile(r"Национальн|Территориальн|Сводн", re.I)


def text_of(html: str) -> str:
    html = re.sub(r"<(script|style)\b.*?</\1>", " ", html, flags=re.I | re.S)
    return " ".join(TAG_RE.sub(" ", html).replace("&nbsp;", " ").split())


def to_float(s: str) -> float:
    return float(re.sub(r"[ \xa0 ]", "", s).replace("−", "-").replace("–", "-").replace(",", "."))


def short_title(t: str) -> str:
    """От последнего «Национальные/Территориальные/Сводные» до конца скобки."""
    starts = [m.start() for m in TITLE_START_RE.finditer(t)]
    return " ".join(t[starts[-1]:].split())[:200]


def parse_rows(html: str, code: str = DEFAULT_CODE) -> list[dict]:
    """Строки таблиц индикаторов с нужным кодом товара."""
    starts = [m.start() for m in ROW_START_RE.finditer(html)]
    out = []
    for n, start in enumerate(starts):
        head = html[start:start + 400]
        if not re.search(rf"""key=["']{re.escape(code)}["']|data-graph=["']{re.escape(code)}["']""", head, re.I):
            continue
        end = starts[n + 1] if n + 1 < len(starts) else start + 4000
        row_html = html[start:min(end, start + 4000)]
        row_text = text_of(row_html)
        if "{{" in row_text:
            log.warning("строка %s заполняется скриптом страницы, значений в HTML нет: %s", code, row_text[:200])
            continue
        after_code = row_text.split(code, 1)[-1]
        val = VALUE_RE.search(after_code)
        if not val:
            log.warning("в строке %s не найдено значение: %s", code, row_text[:200])
            continue
        chg = CHANGE_RE.search(after_code[val.end():])
        name = after_code[:val.start()].strip(" :–-") or code
        before = text_of(html[max(0, start - 6000):start])
        titles = list(TITLE_RE.finditer(before))
        out.append({
            "code": code,
            "name": name,
            "value": to_float(val[1]),
            "change": to_float(chg[1].replace(" ", "")) if chg else None,
            "title": short_title(titles[-1][0]) if titles else "",
            "_text": row_text,
        })
    for i, r in enumerate(out, 1):
        r["table"] = i
    return out


def load_table(path: Path) -> list[dict]:
    if not path.exists():
        return []
    keys = {title: key for key, title in COLUMNS}
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [{keys.get(k, k): v for k, v in r.items()} for r in csv.DictReader(f, delimiter=";")]


def fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return (f"{v:.0f}" if v.is_integer() else f"{v:.2f}").replace(".", ",")
    return str(v)


def num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except ValueError:
        return v


def save(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([t for _, t in COLUMNS])
        for r in rows:
            w.writerow([fmt(r.get(k)) for k, _ in COLUMNS])

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Фенол"
    ws.append([t for _, t in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        line = []
        for k, _ in COLUMNS:
            v = r.get(k)
            if k == "date" and v:
                v = dt.date.fromisoformat(str(v))
            elif k in ("value", "change", "table"):
                v = num(v)
            line.append(v)
        ws.append(line)
    for (c,) in ws.iter_rows(min_row=2, max_col=1):
        c.number_format = "DD.MM.YYYY"
    for (c,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        c.number_format = "# ##0"
    for col, w in zip("ABCDEFGH", (12, 8, 14, 16, 12, 70, 10, 45)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path.with_suffix(".xlsx"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", type=Path, default=DEFAULT_CSV, help="путь к таблице (по умолчанию: %(default)s)")
    ap.add_argument("--url", action="append", help="страница индикаторов (можно несколько раз)")
    ap.add_argument("--code", default=DEFAULT_CODE, help="код товара (по умолчанию: %(default)s)")
    ap.add_argument("--file", type=Path, help="разобрать сохранённую страницу вместо загрузки")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)s %(message)s")

    today = dt.date.today().isoformat()
    collected, errors = [], 0
    pages = [(str(args.file), args.file.read_text(encoding="utf-8"))] if args.file else []
    if not args.file:
        session = requests.Session()
        session.headers.update(HEADERS)
        for url in args.url or DEFAULT_URLS:
            try:
                resp = session.get(url, timeout=60)
                resp.raise_for_status()
                pages.append((url, resp.text))
            except requests.RequestException as e:
                log.error("не удалось открыть %s: %s", url, e)
                errors += 1

    for url, html in pages:
        rows = parse_rows(html, args.code)
        if not rows:
            log.warning("%s: строк %s не найдено", url, args.code)
            errors += 1
        for r in rows:
            r.update(date=today, page=url)
            log.info("%s %s = %s руб./т (%s%%) — %s", r["code"], r["name"], fmt(r["value"]),
                     fmt(r["change"]) if r["change"] is not None else "?", r["title"] or f"таблица {r['table']}")
            log.debug("  текст строки: %s", r.pop("_text"))
            r.pop("_text", None)
            collected.append(r)

    if collected:
        key = lambda r: (str(r.get("date")), str(r.get("page")), str(r.get("table")))
        merged = {key(r): r for r in load_table(args.csv)}
        merged.update({key(r): r for r in collected})
        rows = sorted(merged.values(), key=key)
        save(args.csv, rows)
        log.info("записано строк за %s: %d, всего в таблице: %d (%s)", today, len(collected), len(rows), args.csv)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
